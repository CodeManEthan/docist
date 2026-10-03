"""Office files (design launch-hardening v0.4 §4): the central-directory check,
the archive check in front of every mammoth and openpyxl open, s31.2, and
the spreadsheet cell count. Limits are patched low; fixtures are a few KB.
"""
import io
import re
import time
import zipfile

import openpyxl
import pytest

import converters
import transforms
from converters import ConversionError, docx_converter
from converters.options import RenderOptions
from pdf_ops import limits, office
from transforms import TransformError

import word_fixtures as wf

ARCHIVE = limits.archive_message()


def small_docx(path):
    return wf.write_docx(path, {'word/document.xml': wf.document(wf.para('hello'), '')})


def small_xlsx(path, rows=3):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, rows + 1):
        ws.append([f'r{r}', r])
    wb.save(path)
    return path


def dimension_bomb(path):
    """A 5 KB workbook declaring A1:XFD1048576, with its last row present, so
    openpyxl's read-only mode pads every row to 16,384 cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'a'
    ws['A1048576'] = 'z'
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(buf) as zin:
        parts = {name: zin.read(name) for name in zin.namelist()}
    sheet = parts['xl/worksheets/sheet1.xml'].decode()
    parts['xl/worksheets/sheet1.xml'] = re.sub(
        r'<dimension ref="[^"]*"', '<dimension ref="A1:XFD1048576"', sheet).encode()
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)
    return path


def _office_callables(src):
    """Every registered converter and transform that reads ``src`` files,
    enumerated from the registries (pivots included)."""
    found = []
    convert = converters.get_converter(src)
    if convert is not None:
        found.append((f'converter {src}', convert))
    for target in transforms.targets_for(src):
        found.append((f'transform {src}->{target}', transforms.get_transform(src, target)))
    return found


def _call(func, src, tmp_path, opts=None):
    return func(str(src), str(tmp_path / 'out.bin'), opts or RenderOptions())


# --------------------------------------------------------------------------
# §4.1 the central directory, before zipfile reads it
# --------------------------------------------------------------------------
def test_central_directory_over_the_limit_is_refused(tmp_path, monkeypatch):
    src = small_docx(tmp_path / 'a.docx')
    office.check_central_directory(str(src))     # passes at the real limit
    monkeypatch.setattr(limits, 'MAX_CENTRAL_DIR_BYTES', 64)
    with pytest.raises(office.ArchiveError, match='central directory'):
        office.check_central_directory(str(src))
    with pytest.raises(office.ArchiveError):
        office.check_package(str(src))


def test_central_directory_is_read_before_zipfile_builds_entries(tmp_path, monkeypatch):
    src = small_docx(tmp_path / 'a.docx')
    monkeypatch.setattr(limits, 'MAX_CENTRAL_DIR_BYTES', 64)

    def no_zipfile(*_a, **_k):
        raise AssertionError('zipfile opened the archive')
    monkeypatch.setattr(office.zipfile, 'ZipFile', no_zipfile)
    with pytest.raises(office.ArchiveError):
        office.check_package(str(src))


def test_central_directory_of_a_zip64_archive_is_read(tmp_path, monkeypatch):
    src = tmp_path / 'z64.zip'
    monkeypatch.setattr(zipfile, 'ZIP64_LIMIT', 0)   # force Zip64 end records
    with zipfile.ZipFile(src, 'w', allowZip64=True) as zf:
        for i in range(20):
            zf.writestr(f'part{i}.xml', 'x')
    data = bytearray(src.read_bytes())
    assert b'PK\x06\x06' in data
    # Zero the classic record's directory size: only the Zip64 record tells.
    at = data.rfind(b'PK\x05\x06')
    data[at + 12:at + 16] = b'\0\0\0\0'
    src.write_bytes(bytes(data))
    with zipfile.ZipFile(src) as zf:          # zipfile reads the Zip64 size too
        assert len(zf.infolist()) == 20
    office.check_central_directory(str(src))
    monkeypatch.setattr(limits, 'MAX_CENTRAL_DIR_BYTES', 200)
    with pytest.raises(office.ArchiveError):
        office.check_central_directory(str(src))


def test_a_file_that_isnt_a_zip_is_left_to_zipfile(tmp_path):
    src = tmp_path / 'junk.docx'
    src.write_bytes(b'not a zip at all' * 10)
    office.check_central_directory(str(src))   # no record: no opinion
    with pytest.raises(office.OfficeError):
        office.check_package(str(src))


# --------------------------------------------------------------------------
# §4.2 the archive check in front of every Office open
# --------------------------------------------------------------------------
@pytest.mark.parametrize('src_ext', ['.docx', '.xlsx'])
def test_every_office_reader_runs_the_archive_check(tmp_path, monkeypatch, src_ext):
    src = (small_docx if src_ext == '.docx' else small_xlsx)(tmp_path / f'a{src_ext}')
    monkeypatch.setattr(limits, 'MAX_CENTRAL_DIR_BYTES', 64)

    def opened(*_a, **_k):
        raise AssertionError('a library opened the file')
    for name in ('convert_to_html', 'convert_to_markdown', 'extract_raw_text'):
        monkeypatch.setattr(f'mammoth.{name}', opened)
    monkeypatch.setattr(openpyxl, 'load_workbook', opened)

    callables = _office_callables(src_ext)
    assert len(callables) >= 3
    for label, func in callables:
        for engine in ('reflow', 'libreoffice'):
            with pytest.raises((ConversionError, TransformError)) as info:
                _call(func, src, tmp_path, RenderOptions(word_engine=engine))
            assert ARCHIVE in str(info.value), label
            # The chain carries the limit, so a job reports the clean sentence.
            cause = info.value
            while cause is not None and not isinstance(cause, limits.LimitError):
                cause = cause.__cause__
            assert cause is not None and cause.kind == 'archive', label


def test_a_damaged_xlsx_gets_the_spreadsheet_message(tmp_path):
    src = tmp_path / 'bad.xlsx'
    src.write_bytes(b'PK\x03\x04 this is not really a workbook')
    for label, func in _office_callables('.xlsx'):
        with pytest.raises((ConversionError, TransformError)) as info:
            _call(func, src, tmp_path)
        message = str(info.value)
        assert 'spreadsheet' in message, label
        assert 'Word' not in message and '.docx' not in message, label


def test_a_damaged_docx_gets_the_word_message(tmp_path):
    src = tmp_path / 'bad.docx'
    src.write_bytes(b'PK\x03\x04 this is not really a document')
    for label, func in _office_callables('.docx'):
        with pytest.raises((ConversionError, TransformError)) as info:
            _call(func, src, tmp_path)
        assert 'not a valid .docx file' in str(info.value), label


def test_s31_2_a_paid_file_the_strip_refuses_is_checked_again_in_the_reflow(
        tmp_path, monkeypatch):
    src = small_docx(tmp_path / 'paid.docx')
    calls = []
    real_check = office.check_package

    def spy(path):
        calls.append(path)
        return real_check(path)

    def strip_refuses(*_a, **_k):
        raise office.OfficeError('the strip refused this file')

    monkeypatch.setattr(office, 'check_package', spy)
    monkeypatch.setattr(office, 'docx_to_pdf', strip_refuses)
    opts = RenderOptions(word_engine='libreoffice')
    converters.get_converter('.docx')(str(src), str(tmp_path / 'out.pdf'), opts)
    assert len(calls) == 2            # once before the engine, once in the reflow
    assert docx_converter.FALLBACK_NOTE in opts.notes


def test_s31_2_the_reflows_own_check_refuses(tmp_path, monkeypatch):
    src = small_docx(tmp_path / 'paid.docx')
    monkeypatch.setattr(office, 'docx_to_pdf',
                        lambda *_a, **_k: (_ for _ in ()).throw(office.OfficeError('refused')))
    seen = []
    real_check = office.check_package

    def second_call_refuses(path):
        seen.append(path)
        if len(seen) == 2:
            raise office.ArchiveError('too many parts')
        return real_check(path)
    monkeypatch.setattr(office, 'check_package', second_call_refuses)
    with pytest.raises(ConversionError, match=ARCHIVE):
        converters.get_converter('.docx')(str(src), str(tmp_path / 'out.pdf'),
                                          RenderOptions(word_engine='libreoffice'))


# --------------------------------------------------------------------------
# §4.4 table cells, counted as openpyxl yields them
# --------------------------------------------------------------------------
def test_a_declared_dimension_is_refused_by_the_cell_count_quickly(tmp_path):
    src = dimension_bomb(tmp_path / 'dim.xlsx')
    assert src.stat().st_size < 10_000
    for label, func in _office_callables('.xlsx'):
        started = time.monotonic()
        with pytest.raises((limits.LimitError, ConversionError, TransformError)) as info:
            _call(func, src, tmp_path)
        assert time.monotonic() - started < 5, label
        assert str(info.value) == limits.cells_message(), label


def test_the_cell_count_spans_the_request(tmp_path, monkeypatch):
    monkeypatch.setattr(limits, 'TABLE_CELLS_READ', 10)
    src = small_xlsx(tmp_path / 'a.xlsx', rows=3)      # 6 cells
    opts = RenderOptions()
    converters.get_converter('.xlsx')(str(src), str(tmp_path / 'one.pdf'), opts)
    assert opts.cells_read == 6
    with pytest.raises(limits.LimitError, match='more than 10 cells'):
        converters.get_converter('.xlsx')(str(src), str(tmp_path / 'two.pdf'), opts)


def test_a_real_workbook_still_converts(tmp_path):
    src = small_xlsx(tmp_path / 'a.xlsx', rows=50)
    for label, func in _office_callables('.xlsx'):
        _call(func, src, tmp_path)


def test_route_answers_an_archive_refusal_with_the_sentence(client, tmp_path, monkeypatch):
    src = small_docx(tmp_path / 'a.docx')
    monkeypatch.setattr(limits, 'MAX_CENTRAL_DIR_BYTES', 64)
    for target in ('.pdf', '.txt', '.png'):
        response = client.post('/convert/run', data={
            'file': (io.BytesIO(src.read_bytes()), 'a.docx'), 'target': target,
        }, content_type='multipart/form-data')
        assert response.status_code == 400, target
        assert response.get_json()['error'] == ARCHIVE, target


def test_route_answers_the_cell_count_with_400(client, tmp_path):
    src = dimension_bomb(tmp_path / 'dim.xlsx')
    for target in ('.csv', '.json', '.pdf'):
        response = client.post('/convert/run', data={
            'file': (io.BytesIO(src.read_bytes()), 'dim.xlsx'), 'target': target,
        }, content_type='multipart/form-data')
        assert response.status_code == 400, target
        assert response.get_json()['error'] == limits.cells_message(), target

"""Tests for the spreadsheet (.csv/.xlsx) and RTF (.rtf) converter plugins.

All fixtures are built programmatically in ``tmp_path`` -- no binaries are
committed. Valid outputs are checked with pypdf (page count + extracted text
markers); error cases assert ``ConversionError``; the registry is checked for
the new extensions; and one route-level check uploads a CSV through /upload.
"""
import io

import openpyxl
import pytest
from pypdf import PdfReader

import converters
from converters import (
    ConversionError,
    get_converter,
    spreadsheet_converter,
    supported_extensions,
)


# --------------------------------------------------------------------------
# Markers embedded in fixtures so extracted text can be spot-checked.
# --------------------------------------------------------------------------
CSV_QUOTED = "Paris, FR"          # a quoted field containing a comma
CSV_UNICODE = "café"
CSV_SEMI = "SemiColMarker"

XLSX_SHEET_A = "AlphaSheet"
XLSX_SHEET_B = "BetaSheet"
XLSX_CELL_A = "CellAlphaXi"
XLSX_CELL_B = "CellBetaOmega"
XLSX_WIDE_SHEET = "WideSheet"

RTF_TEXT = "RtfBodyMarkerZeta"


# --------------------------------------------------------------------------
# Programmatic fixture builders
# --------------------------------------------------------------------------
def _build_csv(path):
    """Comma-delimited CSV with a quoted comma-bearing field + unicode."""
    path.write_text(
        "Name,City,Note\n"
        f'Alice,"{CSV_QUOTED}",{CSV_UNICODE}\n'
        "Bob,London,plain\n",
        encoding="utf-8",
    )
    return path


def _build_csv_semicolon(path):
    """Semicolon-delimited CSV (delimiter must be sniffed, not assumed)."""
    path.write_text(
        "Name;City;Tag\n"
        f"Zoe;Berlin;{CSV_SEMI}\n"
        "Yan;Oslo;other\n",
        encoding="utf-8",
    )
    return path


def _build_xlsx(path):
    """Workbook with 2 populated sheets, a wide sheet, and an empty sheet."""
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = XLSX_SHEET_A
    a.append(["Head1", "Head2"])
    a.append([XLSX_CELL_A, 123])

    b = wb.create_sheet(XLSX_SHEET_B)
    b.append(["ColX", "ColY"])
    b.append([XLSX_CELL_B, 4.5])

    wide = wb.create_sheet(XLSX_WIDE_SHEET)
    wide.append([f"C{i}" for i in range(14)])
    wide.append([f"v{i}" for i in range(14)])

    wb.create_sheet("EmptyTab")  # left empty on purpose
    wb.save(path)
    return path


def _build_wide_xlsx(path):
    """A single very wide sheet (forces landscape + column capping)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OnlyWide"
    ws.append([f"Column_{i:02d}" for i in range(20)])
    ws.append(["x" * 300] + [f"val{i}" for i in range(19)])  # extreme cell
    wb.save(path)
    return path


def _build_rtf(path):
    r"""A simple RTF file with bold formatting (formatting is dropped)."""
    path.write_text(
        r"{\rtf1\ansi\deff0 {\b Bold heading}\par "
        + RTF_TEXT
        + r" and some more text.\par}",
        encoding="utf-8",
    )
    return path


def _extract(pdf_path):
    reader = PdfReader(str(pdf_path))
    return len(reader.pages), "\n".join(p.extract_text() or "" for p in reader.pages)


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------
def test_csv_converts_to_valid_pdf_with_content(tmp_path):
    src = _build_csv(tmp_path / "in.csv")
    out = tmp_path / "out.pdf"
    get_converter(".csv")(str(src), str(out))
    pages, text = _extract(out)
    assert pages >= 1
    # Quoted comma-field kept intact, and unicode survived.
    assert "Paris" in text
    assert CSV_UNICODE in text
    assert "Alice" in text and "Bob" in text


def test_csv_semicolon_delimiter_is_sniffed(tmp_path):
    src = _build_csv_semicolon(tmp_path / "semi.csv")
    out = tmp_path / "semi.pdf"
    spreadsheet_converter.convert(str(src), str(out))
    pages, text = _extract(out)
    assert pages >= 1
    assert CSV_SEMI in text
    assert "Berlin" in text


def test_csv_latin1_fallback(tmp_path):
    src = tmp_path / "latin.csv"
    # 0xE9 is 'é' in latin-1 but invalid standalone UTF-8.
    src.write_bytes(b"Name,Note\nJose,caf\xe9 latin\n")
    out = tmp_path / "latin.pdf"
    spreadsheet_converter.convert(str(src), str(out))
    pages, _ = _extract(out)
    assert pages >= 1


def test_empty_csv_yields_one_page_pdf(tmp_path):
    src = tmp_path / "empty.csv"
    src.write_text("", encoding="utf-8")
    out = tmp_path / "empty.pdf"
    spreadsheet_converter.convert(str(src), str(out))
    pages, text = _extract(out)
    assert pages == 1
    assert "empty" in text.lower()


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------
def test_xlsx_multi_sheet_headings_and_cells(tmp_path):
    src = _build_xlsx(tmp_path / "book.xlsx")
    out = tmp_path / "book.pdf"
    get_converter(".xlsx")(str(src), str(out))
    pages, text = _extract(out)
    assert pages >= 1
    # Both sheet names appear as headings, their cells are present.
    assert XLSX_SHEET_A in text
    assert XLSX_SHEET_B in text
    assert XLSX_CELL_A in text
    assert XLSX_CELL_B in text
    # The empty sheet is skipped.
    assert "EmptyTab" not in text


def test_xlsx_wide_table_does_not_error(tmp_path):
    src = _build_wide_xlsx(tmp_path / "wide.xlsx")
    out = tmp_path / "wide.pdf"
    spreadsheet_converter.convert(str(src), str(out))
    pages, text = _extract(out)
    assert pages >= 1
    # A non-extreme column value still shows up.
    assert "val5" in text


def test_corrupt_xlsx_raises(tmp_path):
    src = tmp_path / "broken.xlsx"
    src.write_bytes(b"this is definitely not a real xlsx zip payload")
    out = tmp_path / "broken.pdf"
    with pytest.raises(ConversionError):
        spreadsheet_converter.convert(str(src), str(out))


# --------------------------------------------------------------------------
# RTF
# --------------------------------------------------------------------------
def test_rtf_converts_and_extracts_text(tmp_path):
    src = _build_rtf(tmp_path / "in.rtf")
    out = tmp_path / "out.pdf"
    get_converter(".rtf")(str(src), str(out))
    pages, text = _extract(out)
    assert pages >= 1
    assert RTF_TEXT in text
    assert "Bold heading" in text  # text kept even though bold is dropped


def test_non_rtf_file_raises(tmp_path):
    src = tmp_path / "fake.rtf"
    src.write_text("this is just plain text, not rtf at all", encoding="utf-8")
    out = tmp_path / "fake.pdf"
    with pytest.raises(ConversionError):
        get_converter(".rtf")(str(src), str(out))


def test_garbage_rtf_bytes_raise(tmp_path):
    src = tmp_path / "garbage.rtf"
    src.write_bytes(b"\x00\x01\x02 binary noise " * 20)
    out = tmp_path / "garbage.pdf"
    with pytest.raises(ConversionError):
        spreadsheet_converter  # noqa: F841 - keep import used
        get_converter(".rtf")(str(src), str(out))


# --------------------------------------------------------------------------
# Registry
# --------------------------------------------------------------------------
def test_registry_includes_new_extensions():
    exts = supported_extensions()
    for ext in (".csv", ".xlsx", ".rtf"):
        assert ext in exts
        assert get_converter(ext) is not None


# --------------------------------------------------------------------------
# Route-level check: upload a CSV through POST /upload
# --------------------------------------------------------------------------
def test_upload_csv_merges_successfully(client):
    csv_bytes = (
        "Name,City,Note\n"
        f'Alice,"{CSV_QUOTED}",{CSV_UNICODE}\n'
        "Bob,London,plain\n"
    ).encode("utf-8")

    data = {"files[]": [(io.BytesIO(csv_bytes), "data.csv")]}
    resp = client.post("/upload", data=data, content_type="multipart/form-data")
    assert resp.status_code == 200, resp.data
    payload = resp.get_json()
    assert payload["success"] is True

    filename = payload["filename"]
    dl = client.get(f"/download?filename={filename}")
    assert dl.status_code == 200, dl.data
    reader = PdfReader(io.BytesIO(dl.data))
    # A single short CSV -> 1 content page, padded to an even count (2) by
    # the app's blank-page logic. Sanity-check it produced pages.
    assert len(reader.pages) >= 1

"""Transform plugin: data-format conversions (CSV / XLSX / JSON / YAML / HTML).

Registered pairs (see ``TRANSFORMS`` at the bottom)::

    .csv  -> .xlsx      .xlsx -> .csv
    .csv  -> .json      .json -> .csv
    .xlsx -> .json
    .json <-> .yaml     .json <-> .yml   (.yml is a straight alias for .yaml)
    .csv  -> .html

Each function has the signature ``func(input_path, output_path)`` and writes
the result to ``output_path`` (returning that path). Failures raise
``transforms.TransformError`` with a human-readable message.

Value-typing policy (applied to CSV cells and, after stringification, to XLSX
cells)
    Values are kept as **text** by default. A value is promoted to a *number*
    only when its trimmed text is a clean integer or float literal:

    * Integer:  ``^[+-]?(0|[1-9][0-9]*)$`` -> ``int``.  Note the deliberate
      rejection of leading zeros ("007", "0123"): such strings are almost
      always identifiers (zip codes, phone numbers, part numbers) whose zeros
      must be preserved, so they stay text. Plain ``0`` / ``-3`` / ``+42`` do
      convert.
    * Float: must contain a decimal point or an exponent
      (e.g. ``1.5``, ``.5``, ``1.``, ``1e3``, ``1.5e-3``) -> ``float``.
      ``inf`` / ``nan`` / underscore or hex forms never match, so they stay
      text.

    Empty cells become ``null`` in JSON and empty cells in XLSX/CSV.

Header policy (CSV/XLSX -> JSON, the object keys)
    The first row is the header. Keys are de-duplicated left-to-right: an empty
    header cell becomes ``col_N`` (1-based column index); a repeated key gets a
    ``_2`` / ``_3`` ... suffix on its 2nd, 3rd, ... occurrence.

CSV reading (everywhere)
    Delimiter is sniffed with :class:`csv.Sniffer` (``,;\\t|``), falling back to
    a comma. Bytes decode as UTF-8 with a latin-1 fallback. Ragged rows are
    right-padded with empty strings to a common width.
"""
import csv
import html
import io
import json
import re

import openpyxl
import yaml

from . import TransformError


# --------------------------------------------------------------------------
# Value typing
# --------------------------------------------------------------------------
_INT_RE = re.compile(r'^[+-]?(?:0|[1-9][0-9]*)$')
_FLOAT_RE = re.compile(
    r'''^[+-]?
        (?:
            (?:\d+\.\d* | \.\d+) (?:[eE][+-]?\d+)?   # requires a decimal point
          | \d+ [eE][+-]?\d+                          # int mantissa + exponent
        )$''',
    re.VERBOSE,
)


def _coerce_scalar(text):
    """Return ``int``/``float`` for a clean numeric literal, else the string.

    See the module-level typing policy. Non-strings are returned unchanged.
    """
    if not isinstance(text, str):
        return text
    s = text.strip()
    if s == '':
        return text
    if _INT_RE.match(s):
        try:
            return int(s)
        except ValueError:
            return text
    if _FLOAT_RE.match(s):
        try:
            return float(s)
        except ValueError:
            return text
    return text


def _stringify(value):
    """Turn a cell/JSON scalar into CSV text (``None`` -> ``''``)."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# --------------------------------------------------------------------------
# CSV reading / header + record helpers
# --------------------------------------------------------------------------
def _read_csv_rows(input_path):
    """Read a CSV file into a list of equal-width string rows.

    Sniffs the delimiter (comma fallback), decodes UTF-8 then latin-1, and
    right-pads ragged rows. Returns ``[]`` for an empty/blank file.
    """
    with open(input_path, 'rb') as fh:
        raw = fh.read()
    if not raw.strip():
        return []
    try:
        text = raw.decode('utf-8')
    except UnicodeDecodeError:
        text = raw.decode('latin-1', errors='replace')

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel  # comma-delimited default

    reader = csv.reader(io.StringIO(text), dialect)
    rows = [list(row) for row in reader]
    if rows:
        width = max(len(r) for r in rows)
        for r in rows:
            r.extend([''] * (width - len(r)))
    return rows


def _dedupe_headers(header_cells):
    """Map raw header cells to unique keys per the header policy."""
    used = {}
    keys = []
    for i, cell in enumerate(header_cells):
        name = (cell or '').strip()
        if name == '':
            name = f'col_{i + 1}'
        if name in used:
            used[name] += 1
            candidate = f'{name}_{used[name]}'
            while candidate in used:
                used[name] += 1
                candidate = f'{name}_{used[name]}'
            used[candidate] = 1
            keys.append(candidate)
        else:
            used[name] = 1
            keys.append(name)
    return keys


def _rows_to_records(rows):
    """Turn ``[header, *data]`` string rows into a list of typed dicts."""
    if not rows:
        return []
    keys = _dedupe_headers(rows[0])
    records = []
    for row in rows[1:]:
        obj = {}
        for c, key in enumerate(keys):
            cell = row[c] if c < len(row) else ''
            obj[key] = None if cell.strip() == '' else _coerce_scalar(cell)
        records.append(obj)
    return records


# --------------------------------------------------------------------------
# XLSX reading
# --------------------------------------------------------------------------
def _open_workbook(input_path):
    try:
        return openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any openpyxl/zip failure
        raise TransformError(
            f"Could not open spreadsheet '{input_path}': {exc}"
        ) from exc


def _sheet_rows(ws):
    """Read a worksheet into stringified rows, dropping trailing empty rows."""
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([_stringify(v) for v in row])
    while rows and all(c == '' for c in rows[-1]):
        rows.pop()
    return rows


# --------------------------------------------------------------------------
# JSON reading
# --------------------------------------------------------------------------
def _load_json(input_path):
    try:
        with open(input_path, encoding='utf-8') as fh:
            return json.load(fh)
    except (json.JSONDecodeError, UnicodeDecodeError, OSError) as exc:
        raise TransformError(f'Invalid JSON input: {exc}') from exc


def _dump_json(data, output_path):
    with open(output_path, 'w', encoding='utf-8') as fh:
        json.dump(data, fh, indent=2, ensure_ascii=False)
        fh.write('\n')


# --------------------------------------------------------------------------
# Transforms
# --------------------------------------------------------------------------
def csv_to_xlsx(input_path, output_path):
    """CSV -> XLSX workbook, single sheet named ``Sheet1``.

    Cells are written as text except clean int/float literals, which become
    numbers (typing policy). Empty cells stay empty.
    """
    rows = _read_csv_rows(input_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    for row in rows:
        ws.append([None if c.strip() == '' else _coerce_scalar(c) for c in row])
    wb.save(output_path)
    return output_path


def xlsx_to_csv(input_path, output_path):
    """XLSX -> CSV (comma-delimited, minimal quoting).

    Only the **first** worksheet is exported. For multi-sheet workbooks the
    remaining sheets are intentionally dropped (CSV is single-table); use the
    ``.xlsx -> .json`` transform to capture every sheet.
    """
    wb = _open_workbook(input_path)
    try:
        ws = wb.worksheets[0]
        rows = _sheet_rows(ws)
    finally:
        wb.close()
    with open(output_path, 'w', newline='', encoding='utf-8') as fh:
        writer = csv.writer(fh, quoting=csv.QUOTE_MINIMAL)
        for row in rows:
            writer.writerow(row)
    return output_path


def csv_to_json(input_path, output_path):
    """CSV -> JSON array of objects (header row -> keys).

    Values follow the typing policy; empty cells become ``null``. UTF-8,
    ``indent=2``.
    """
    records = _rows_to_records(_read_csv_rows(input_path))
    _dump_json(records, output_path)
    return output_path


def json_to_csv(input_path, output_path):
    """JSON -> CSV.

    Accepts a top-level array of *flat* objects (union of keys, first-seen
    order, becomes the header; missing keys -> empty cell) OR an array of
    scalars (single ``value`` column). Any nested value (a dict/list inside an
    object), a non-array top level, or a mixed/other shape raises
    ``TransformError``.
    """
    data = _load_json(input_path)
    if not isinstance(data, list):
        raise TransformError(
            'JSON -> CSV expects a top-level array, got '
            f'{type(data).__name__}.'
        )

    with open(output_path, 'w', newline='', encoding='utf-8') as fh:
        writer = csv.writer(fh, quoting=csv.QUOTE_MINIMAL)
        if not data:
            return output_path  # empty array -> empty file

        if all(not isinstance(item, (dict, list)) for item in data):
            # Array of scalars -> single "value" column.
            writer.writerow(['value'])
            for item in data:
                writer.writerow([_stringify(item)])
            return output_path

        if all(isinstance(item, dict) for item in data):
            headers = []
            seen = set()
            for item in data:
                for value in item.values():
                    if isinstance(value, (dict, list)):
                        raise TransformError(
                            'JSON -> CSV cannot flatten nested objects/arrays; '
                            'every object value must be a scalar.'
                        )
                for key in item.keys():
                    if key not in seen:
                        seen.add(key)
                        headers.append(key)
            writer.writerow(headers)
            for item in data:
                writer.writerow([_stringify(item.get(k)) for k in headers])
            return output_path

    raise TransformError(
        'JSON -> CSV expects an array of flat objects or an array of scalars; '
        'got a mixed or unsupported array shape.'
    )


def xlsx_to_json(input_path, output_path):
    """XLSX -> JSON object ``{sheet_name: [objects]}`` for every sheet.

    Each sheet uses the same header/typing rules as ``.csv -> .json``. Empty
    sheets map to an empty array.
    """
    wb = _open_workbook(input_path)
    try:
        result = {ws.title: _rows_to_records(_sheet_rows(ws))
                  for ws in wb.worksheets}
    finally:
        wb.close()
    _dump_json(result, output_path)
    return output_path


def json_to_yaml(input_path, output_path):
    """JSON -> YAML (block style, unicode preserved, key order kept)."""
    data = _load_json(input_path)
    with open(output_path, 'w', encoding='utf-8') as fh:
        yaml.safe_dump(
            data, fh,
            default_flow_style=False,
            allow_unicode=True,
            sort_keys=False,
        )
    return output_path


def yaml_to_json(input_path, output_path):
    """YAML -> JSON (UTF-8, ``indent=2``). Invalid YAML -> ``TransformError``."""
    try:
        with open(input_path, encoding='utf-8') as fh:
            data = yaml.safe_load(fh)
    except (yaml.YAMLError, UnicodeDecodeError, OSError) as exc:
        raise TransformError(f'Invalid YAML input: {exc}') from exc
    _dump_json(data, output_path)
    return output_path


_HTML_CSS = """
  body { font-family: system-ui, sans-serif; margin: 1.5rem; color: #222; }
  table { border-collapse: collapse; font-size: 14px; }
  th, td { border: 1px solid #d0d0d0; padding: 6px 10px; text-align: left; }
  th { background: #f2f5fa; font-weight: 600; }
  tbody tr:nth-child(even) td { background: #fafafa; }
""".strip('\n')


def csv_to_html(input_path, output_path):
    """CSV -> standalone HTML document with a styled table.

    The first row renders as ``<th>`` header cells; every cell value is
    HTML-escaped so embedded markup (e.g. ``<script>``) is inert.
    """
    rows = _read_csv_rows(input_path)
    parts = [
        '<!DOCTYPE html>',
        '<html lang="en">',
        '<head>',
        '<meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        '<title>CSV Table</title>',
        f'<style>\n{_HTML_CSS}\n</style>',
        '</head>',
        '<body>',
        '<table>',
    ]
    if rows:
        header, *body = rows
        parts.append('<thead><tr>')
        parts.extend(f'<th>{html.escape(c)}</th>' for c in header)
        parts.append('</tr></thead>')
        parts.append('<tbody>')
        for row in body:
            parts.append('<tr>')
            parts.extend(f'<td>{html.escape(c)}</td>' for c in row)
            parts.append('</tr>')
        parts.append('</tbody>')
    parts += ['</table>', '</body>', '</html>']
    with open(output_path, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(parts) + '\n')
    return output_path


TRANSFORMS = {
    ('.csv', '.xlsx'): csv_to_xlsx,
    ('.xlsx', '.csv'): xlsx_to_csv,
    ('.csv', '.json'): csv_to_json,
    ('.json', '.csv'): json_to_csv,
    ('.xlsx', '.json'): xlsx_to_json,
    ('.json', '.yaml'): json_to_yaml,
    ('.yaml', '.json'): yaml_to_json,
    ('.json', '.yml'): json_to_yaml,   # .yml is an alias for .yaml
    ('.yml', '.json'): yaml_to_json,
    ('.csv', '.html'): csv_to_html,
}
